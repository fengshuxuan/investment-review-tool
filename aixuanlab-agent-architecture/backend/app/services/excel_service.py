import uuid
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.core.config import get_settings
from app.schemas.ai import ExcelGenerateResponse, PracticeQuestion


class ExcelService:
    """Generate and modify Excel files.

    MVP implements practice-question-to-xlsx. Later extend with workbook analysis and safe operations.
    """

    def __init__(self) -> None:
        self.settings = get_settings()
        Path(self.settings.generated_dir).mkdir(parents=True, exist_ok=True)

    def generate_practice_workbook(self, user_id: str, practice: PracticeQuestion, include_answer_sheet: bool = True) -> ExcelGenerateResponse:
        file_id = str(uuid.uuid4())
        user_dir = Path(self.settings.generated_dir) / user_id
        user_dir.mkdir(parents=True, exist_ok=True)
        file_path = user_dir / f"{file_id}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "练习题"

        ws["A1"] = practice.title
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"场景：{practice.scenario or ''}"
        ws["A3"] = f"目标函数：{', '.join(practice.function_tags)}"
        ws["A4"] = f"请在 {practice.target_cell} 输入公式完成练习。"

        start_row = 6
        for col_idx, header in enumerate(practice.headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E2F0D9")

        for row_idx, row in enumerate(practice.rows, start=start_row + 1):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(max_length + 2, 12), 30)

        if include_answer_sheet:
            answer_ws = wb.create_sheet("参考答案")
            answer_ws["A1"] = "目标单元格"
            answer_ws["B1"] = "参考公式"
            answer_ws["A2"] = practice.target_cell
            answer_ws["B2"] = practice.answer_formula
            answer_ws["A4"] = "提示"
            for idx, hint in enumerate(practice.hints, start=5):
                answer_ws.cell(row=idx, column=1, value=hint)
            answer_ws.column_dimensions["A"].width = 20
            answer_ws.column_dimensions["B"].width = 60

        wb.save(file_path)
        return ExcelGenerateResponse(
            file_id=file_id,
            download_url=f"{self.settings.public_file_base_url}/{user_id}/{file_id}.xlsx",
        )
